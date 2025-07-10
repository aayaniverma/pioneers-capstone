import uuid
import os
import shutil
import json
import traceback
from fastapi import APIRouter, UploadFile, File
from fastapi.responses import JSONResponse
from docx import Document
from dotenv import load_dotenv
from openai import OpenAI

# Load environment variables
load_dotenv()

router = APIRouter()

# Folder setup
UPLOAD_FOLDER = "./output_contracts"
VERIFIED_FOLDER = "./verified_con"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(VERIFIED_FOLDER, exist_ok=True)

# Initialize OpenAI client with API key from .env
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# Extract text from .docx
def extract_text_from_docx(docx_path):
    doc = Document(docx_path)
    return "\n".join(para.text.strip() for para in doc.paragraphs if para.text.strip())

# GPT-4 Contract Validation
def ai_verify_contract(contract_text):
    prompt = (
    "You are a legal contract structure validator. Your job is to verify whether specific fields are structurally present in a given M&A contract — regardless of whether the actual data is filled in.\n\n"

    "Your ONLY goal is to confirm the **presence of the field reference**, not the value.\n\n"

    "✅ A field is VALID if:\n"
    "- The field appears in any form, including:\n"
    "   • A placeholder like `{{ Total_Amount }}`\n"
    "   • A label like `Name: ________` or `Title:` even if blank\n"
    "   • A clearly associated sentence (e.g., 'The Seller shall deliver the Assets on the Closing Date')\n"
    "- It may use alternate terms (e.g., 'Buyer' = 'Acquirer')\n\n"

    "❌ A field is INVALID only if:\n"
    "- There is NO mention of the field at all in any form — no label, no reference, no placeholder\n\n"

    "Examples:\n"
    "- 'Payment: {{ Total_Amount }}' ✅ valid (placeholder)\n"
    "- 'Title: __________' ✅ valid (labeled)\n"
    "- 'Payment shall be delivered on the date specified herein.' ✅ valid\n"
    "- No mention of governing law anywhere ❌ invalid\n\n"

    "Now check this contract for the following fields:\n"
    "1. Acquirer Name\n"
    "2. Acquirer Address\n"
    "3. Seller Name\n"
    "4. Seller Address\n"
    "5. Jurisdiction of Acquirer\n"
    "6. Jurisdiction of Seller\n"
    "7. Effective Date\n"
    "8. Closing Date\n"
    "9. Governing Law\n"
    "10. Payment Method or Purchase Price\n"
    "11. Signatory Name A\n"
    "12. Signatory Title A\n"
    "13. Signatory Name B\n"
    "14. Signatory Title B\n\n"

    "You MUST assume all placeholders or labels (even blank) are valid field references.\n\n"

    "Respond ONLY in this JSON format:\n"
    "{\n"
    "  \"status\": \"Verified\" or \"Not Verified\",\n"
    "  \"fields\": {\n"
    "    \"Acquirer Name\": {\"valid\": true/false, \"issue\": \"...\", \"suggestion\": \"...\"},\n"
    "    ... (14 fields total) ...\n"
    "  },\n"
    "  \"summary\": \"One-line summary stating if all fields are structurally present\"\n"
    "}\n\n"

    "Now analyze this contract and return only the JSON:\n\n"
    f"{contract_text}"
)

    try:
        response = client.chat.completions.create(
            model="gpt-4",
            temperature=0,
            messages=[
                {"role": "system", "content": "You are a legal compliance assistant."},
                {"role": "user", "content": prompt}
            ]
        )

        reply = response.choices[0].message.content.strip()
        print("🔍 GPT RAW REPLY:\n", reply)

        # Try to parse response
        try:
            result = json.loads(reply)
            fields = result.get("fields", {})
            invalid_fields = [k for k, v in fields.items() if not v.get("valid", False)]

            if invalid_fields:
                result["status"] = "Not Verified"
                result["summary"] = f"Missing or invalid fields: {', '.join(invalid_fields)}"
            else:
                result["status"] = "Verified"
                result["summary"] = "All required fields are valid."

            return result

        except json.JSONDecodeError:
            return {
                "status": "Not Verified",
                "summary": "OpenAI returned invalid JSON.",
                "fields": {},
                "raw": reply
            }

    except Exception as e:
        return {
            "status": "Not Verified",
            "summary": "OpenAI API failed",
            "fields": {},
            "error": str(e)
        }

# Upload Endpoint
@router.post("/verify-contract")
async def verify_contract(file: UploadFile = File(...)):
    unique_name = f"{uuid.uuid4().hex}_{file.filename}"
    upload_path = os.path.join(UPLOAD_FOLDER, unique_name)

    try:
        contents = await file.read()
        with open(upload_path, "wb") as buffer:
            buffer.write(contents)
        file.file.close()

        # Extract text
        contract_text = extract_text_from_docx(upload_path)

        # Run GPT analysis
        ai_result = ai_verify_contract(contract_text)

        if ai_result.get("status") == "Verified":
            verified_path = os.path.join(VERIFIED_FOLDER, f"verified_{file.filename}")
            shutil.move(upload_path, verified_path)
            return {
                "success": True,
                "verification_status": "Verified",
                "message": ai_result.get("summary", ""),
                "verified_file": f"verified_{file.filename}"
            }

        # Not Verified
        os.remove(upload_path)
        return JSONResponse(status_code=200, content={
            "success": False,
            "verification_status": "Not Verified",
            "field_level_issues": ai_result.get("fields", {}),
            "summary": ai_result.get("summary", ""),
            "raw_output": ai_result
        })

    except Exception as e:
        traceback.print_exc()
        return JSONResponse(status_code=500, content={
            "success": False,
            "errors": [str(e)]
        })
